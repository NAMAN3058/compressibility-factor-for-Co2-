#!/usr/bin/env python3
"""
CO2 absorption dashboard (Peace Software Z values -> absorbed CO2 -> Excel).

Run:   python app.py        (opens http://127.0.0.1:5057 in your browser)

What it does
  1. You upload an Excel workbook (one or many sheets, e.g. 100 psi ... 500 psi, or NaCl / MgCl2 / CaCl2).
  2. Per sheet it finds the Time and pressure columns and the pressure unit (psi / bar / kPa / MPa).
  3. Z1 (first pressure of the sheet, looked up once) and Z2 (every row) come from the Peace Software
     online CO2 calculator - NOT CoolProp - at the temperature you choose.
  4. From Z1/Z2 it calculates Pi/Pf (Pa), absorbed mol, absorbed CO2 (g), g/L, concentration, mol/kg,
     error and minutes, all with settings you can change in the dashboard.
  5. You watch the values fill in live, then download <name>_Z.xlsx. Original sheets and cells are kept;
     a "Calculation Notes" sheet documents every formula, unit and assumption.

Reliability: every distinct pressure is requested once (cache), failed requests are retried, failed rows
are listed with their error message and left blank (never estimated), the workbook is marked INCOMPLETE
if anything failed, and "Retry failed rows only" repeats just the failed lookups.
"""
import io
import os
import re
import sys
import json
import time
import uuid
import threading
import traceback
import webbrowser
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import urljoin

import openpyxl
import requests
from bs4 import BeautifulSoup
from flask import Flask, request, jsonify, send_file, Response

# --------------------------------------------------------------------------- config
SITE_URLS = [u for u in [
    os.environ.get("PEACE_URL"),
    "https://www.peacesoftware.de/einigewerte/co2_e.html",
    "http://www.peacesoftware.de/einigewerte/co2_e.html",
    "https://peace-software.de/einigewerte/co2_e.html",
    "http://peace-software.de/einigewerte/co2_e.html",
] if u]
PORT = int(os.environ.get("PORT", "5057"))
OUT_DIR = Path(__file__).parent / "outputs"
OUT_DIR.mkdir(exist_ok=True)
DEBUG_HTML = Path(__file__).parent / "debug_last_page.html"
UA = {"User-Agent": "Mozilla/5.0 (compatible; ZDashboard/1.0)"}

# factors to bar, used only if the website has no option for the unit you need
TO_BAR = {"psi": 0.06894757293168, "bar": 1.0, "kpa": 0.01, "mpa": 10.0, "atm": 1.01325, "pa": 1e-5}
UNIT_LABEL = {"psi": "psi", "bar": "bar", "kpa": "kPa", "mpa": "MPa", "atm": "atm", "pa": "Pa"}
ALIASES = {"psi": ["psi", "psia", "psig"], "bar": ["bar", "bara"], "kpa": ["kpa"], "mpa": ["mpa"],
           "atm": ["atm"], "pa": ["pa"]}


def norm(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def unit_from_text(s):
    s = str(s or "").lower()
    for key, pat in (("kpa", r"kpa"), ("mpa", r"mpa"), ("psi", r"psi"), ("bar", r"bar"),
                     ("atm", r"\batm\b|atm$"), ("pa", r"(?<![a-z])pa(?![a-z])")):
        if re.search(pat, s):
            return key
    return None


# --------------------------------------------------------------------------- website client
NUM = r"([-+]?(?:\d+[.,]?\d*|[.,]\d+)(?:[eE][-+]?\d+)?)"


def to_float(s):
    s = str(s).strip().replace(" ", "")
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    return float(s)


def _prev_text(el, n=1):
    parts = []
    for s in el.find_all_previous(string=True, limit=15):
        t = s.strip()
        if not t or s.parent.name in ("script", "style", "option", "title"):
            continue
        parts.append(t)
        if len(parts) >= n:
            break
    return " ".join(reversed(parts)).lower()


def parse_z(html, extra_values=None):
    """Find the compressibility factor Z in a result page. Returns float or None."""
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style"]):
        t.decompose()
    cands = []
    for el in soup.find_all(["input", "textarea"]):
        ctx = ((el.get("name") or "") + " " + (el.get("id") or "") + " " + _prev_text(el)).lower()
        cands.append((ctx, el.get("value") or el.get_text()))
    cands += list(extra_values or [])
    for ctx, val in cands:
        ctx = (ctx or "").lower()
        if re.search(r"compress|z[-_ ]?fact|zfactor|(^|[^a-z])z($|[^a-z])", ctx or ""):
            try:
                v = to_float(val)
                if 0 < v < 10:
                    return v
            except Exception:
                pass
    text = soup.get_text(" ", strip=True)
    for pat in (r"compressib[a-z]*(?:\s+factor)?(?:\s*\(?\s*z\s*\)?)?[^0-9+\-.,]{0,20}" + NUM,
                r"(?<![A-Za-z])Z\s*[=:]\s*" + NUM):
        for m in re.finditer(pat, text, re.I):
            try:
                v = to_float(m.group(1))
                if 0 < v < 10:
                    return v
            except Exception:
                continue
    return None


NON_TEXT_TYPES = {"hidden", "submit", "button", "reset", "image", "file", "checkbox", "radio",
                  "password", "color", "range"}
T_LABELS = {"c", "k", "f", "degc", "degk", "degf", "kelvin", "celsius", "fahrenheit"}


def fmtnum(v):
    return format(float(v), ".10g")


def elements_summary(html, limit=30):
    """Compact list of the page's input/select/button elements (for error messages)."""
    soup = BeautifulSoup(html, "html.parser")
    out = []
    for el in soup.find_all(["input", "select", "button", "textarea"]):
        d = [el.name]
        for a in ("type", "name", "id", "value", "onclick", "onchange"):
            v = el.get(a)
            if v:
                d.append(f"{a}={str(v)[:30]}")
        out.append(" ".join(d))
        if len(out) >= limit:
            break
    return "; ".join(out) or "(no input/select/button elements at all)"


class FormSpec:
    """What we learned about the calculator (fields, unit controls, method...)."""

    def __init__(self):
        self.action = None
        self.method = "get"
        self.base = {}
        self.t_name = self.p_name = None      # keys (name / id / position)
        self.t_css = self.p_css = None        # selectors for the browser
        self.t_sel = self.p_sel = None        # (key, [(label, value)], css, kind)
        self.t_fixed = None                   # 'c'/'k' when the unit is only printed as text
        self.p_fixed = None
        self.submit = None
        self.guessed = False
        self.formless = False
        self.page_url = None

    def describe(self):
        return {"action": self.action, "method": self.method, "has_html_form": not self.formless,
                "temperature_field": self.t_name, "pressure_field": self.p_name,
                "temperature_units": [o[0] for o in self.t_sel[1]] if self.t_sel else f"fixed: {self.t_fixed}",
                "pressure_units": [o[0] for o in self.p_sel[1]] if self.p_sel else f"fixed: {self.p_fixed}",
                "fields_guessed_by_order": self.guessed}


def discover(html, page_url):
    soup = BeautifulSoup(html, "html.parser")
    ipos = {id(el): k for k, el in enumerate(soup.find_all("input"))}
    spos = {id(el): k for k, el in enumerate(soup.find_all("select"))}

    def ikey(el):
        return el.get("name") or el.get("id") or f"input#{ipos[id(el)]}"

    def icss(el):
        if el.get("name"):
            return f"[name='{el['name']}']"
        if el.get("id"):
            return f"#{el['id']}"
        return f"xpath=(//input)[{ipos[id(el)] + 1}]"

    def near(el):
        nxt = el.find_next(string=lambda s: s and s.strip() and s.parent.name not in ("script", "style"))
        return (_prev_text(el, 2) + " " + (nxt or "")).lower()

    containers = [(f, False) for f in soup.find_all("form")] + [(soup, True)]
    for cont, formless in containers:
        texts = [i for i in cont.find_all("input")
                 if (i.get("type") or "text").lower() not in NON_TEXT_TYPES and (formless or i.get("name"))]
        if len(texts) < 2:
            continue
        spec = FormSpec()
        spec.page_url = page_url
        spec.formless = formless
        spec.action = urljoin(page_url, (cont.get("action") or "") if not formless else "")
        spec.method = "get" if formless else (cont.get("method") or "get").lower()
        if not formless:
            for el in cont.find_all(["input", "select", "textarea"]):
                name = el.get("name")
                if not name:
                    continue
                typ = (el.get("type") or "text").lower()
                if el.name == "select":
                    opts = el.find_all("option")
                    sel = next((o for o in opts if o.has_attr("selected")), opts[0] if opts else None)
                    spec.base[name] = (sel.get("value", sel.get_text(strip=True)) if sel else "")
                elif typ in ("checkbox", "radio"):
                    if el.has_attr("checked"):
                        spec.base[name] = el.get("value", "on")
                elif typ in ("submit", "button", "reset", "image", "file"):
                    if typ == "submit" and spec.submit is None:
                        spec.submit = (name, el.get("value", ""))
                else:
                    spec.base[name] = el.get("value", "")
        t = p = None
        for use_label in (False, True):
            for i in texts:
                ctx = (ikey(i) + " " + (i.get("id") or "")).lower()
                if use_label:
                    ctx += " " + _prev_text(i).lower()
                if t is None and i is not p and re.search(r"temp|(^|[^a-z])t($|[^a-z])", ctx):
                    t = i
                elif p is None and i is not t and re.search(r"press|druck|(^|[^a-z])p($|[^a-z])", ctx):
                    p = i
        if t is None or p is None or t is p:
            t, p = texts[0], texts[1]
            spec.guessed = True
        spec.t_name, spec.p_name = ikey(t), ikey(p)
        spec.t_css, spec.p_css = icss(t), icss(p)

        def assign(group, last):
            labels = [norm(o[0]) for o in group[1]]
            is_t = any(x in T_LABELS for x in labels)
            is_p = any(unit_from_text(o[0]) for o in group[1])
            if is_p and not is_t and spec.p_sel is None:
                spec.p_sel = group
            elif is_t and spec.t_sel is None:
                spec.t_sel = group
            elif last == "p" and spec.p_sel is None and len(group[1]) > 1:
                spec.p_sel = group
            elif last == "t" and spec.t_sel is None and len(group[1]) > 1:
                spec.t_sel = group

        last = None
        radios = {}
        for el in cont.find_all(["input", "select"]):
            if el.name == "input":
                typ = (el.get("type") or "text").lower()
                if el is t:
                    last = "t"
                elif el is p:
                    last = "p"
                elif typ == "radio" and el.get("name"):
                    lab = el.find_next(string=lambda s: s and s.strip())
                    radios.setdefault(el["name"], []).append(
                        ((lab or el.get("value") or "").strip(), el.get("value") or (lab or "").strip()))
            else:
                key = el.get("name") or el.get("id") or f"select#{spos[id(el)]}"
                css = (f"[name='{el['name']}']" if el.get("name") else f"#{el['id']}" if el.get("id")
                       else f"xpath=(//select)[{spos[id(el)] + 1}]")
                opts = [(o.get_text(strip=True), o.get("value", o.get_text(strip=True)))
                        for o in el.find_all("option")]
                assign((key, opts, css, "select"), last)
        for name, opts in radios.items():
            assign((name, opts, f"input[name='{name}']", "radio"), None)
        if spec.p_sel is None:
            spec.p_fixed = unit_from_text(near(p))
        if spec.t_sel is None:
            spec.t_fixed = "k" if re.search(r"(^|[^a-z])k($|[^a-z])|kelvin", near(t)) else "c"
        return spec
    return None


def pick_option(sel, key):
    """Find the option value of a unit control for a unit key."""
    if not sel:
        return None
    for label, value in sel[1]:
        for n in (norm(label), norm(value)):
            if key == "c" and n in ("c", "degc", "celsius"):
                return value
            if key == "k" and n in ("k", "kelvin", "degk"):
                return value
            if key not in ("c", "k") and n in ALIASES[key]:
                return value
    return None


class SiteError(Exception):
    pass


class ParseFail(SiteError):
    pass


def plan_units(sp, p, p_unit, t, t_unit):
    """Return (p, pressure_option, t, temperature_option). Converts only if the site lacks the unit."""
    popt = topt = None
    if sp.p_sel:
        popt = pick_option(sp.p_sel, p_unit)
        if popt is None:
            popt = pick_option(sp.p_sel, "bar")
            if popt is None:
                raise SiteError(f"Website has no option for {UNIT_LABEL[p_unit]} or bar")
            p = p * TO_BAR[p_unit]
    else:
        if sp.p_fixed is None:
            raise SiteError("Could not find the pressure-unit control on the page - refusing to guess the unit")
        p = p * TO_BAR[p_unit] / TO_BAR[sp.p_fixed]
    if sp.t_sel:
        topt = pick_option(sp.t_sel, t_unit)
        if topt is None:
            other = "k" if t_unit == "c" else "c"
            topt = pick_option(sp.t_sel, other)
            if topt is None:
                raise SiteError("Website has no °C or K option")
            t = t + 273.15 if other == "k" else t - 273.15
    elif (sp.t_fixed or "c") != t_unit:
        t = t + 273.15 if t_unit == "c" else t - 273.15
    return p, popt, t, topt


class RequestsBackend:
    name = "requests"

    def __init__(self, log):
        self.log = log
        self.spec = None
        self.local = threading.local()
        diag = []
        for url in SITE_URLS:
            try:
                r = requests.get(url, headers=UA, timeout=20)
                html = r.content.decode(r.encoding or r.apparent_encoding or "utf-8", "replace")
                DEBUG_HTML.write_text(html, encoding="utf-8", errors="replace")
                soup = BeautifulSoup(html, "html.parser")
                title = soup.title.get_text(strip=True) if soup.title else ""
                diag.append(f"{url} -> {r.url}, status {r.status_code}, title '{title}', "
                            f"forms {len(soup.find_all('form'))}, inputs {len(soup.find_all('input'))}, "
                            f"iframes {len(soup.find_all(['iframe', 'frame']))}, "
                            f"elements: {elements_summary(html, 12)}")
                r.raise_for_status()
                spec = discover(html, r.url)
                if spec and not spec.formless:
                    self.spec = spec
                    self.log(f"Website reached: {r.url}")
                    return
                if spec:
                    diag.append("calculator inputs found but there is no HTML form (JavaScript page)")
            except Exception as e:  # noqa
                diag.append(f"{url} -> error: {e}")
        raise SiteError("No HTML calculator form. " + " || ".join(diag))

    def _session(self):
        s = getattr(self.local, "s", None)
        if s is None:
            s = requests.Session()
            s.headers.update(UA)
            self.local.s = s
        return s

    def z(self, p, p_unit, t, t_unit):
        sp = self.spec
        p, popt, t, topt = plan_units(sp, p, p_unit, t, t_unit)
        payload = dict(sp.base)
        if popt is not None:
            payload[sp.p_sel[0]] = popt
        if topt is not None:
            payload[sp.t_sel[0]] = topt
        payload[sp.p_name] = fmtnum(p)
        payload[sp.t_name] = fmtnum(t)
        if sp.submit and sp.submit[0]:
            payload[sp.submit[0]] = sp.submit[1]
        for attempt in range(3):
            try:
                kw = {"params": payload} if sp.method == "get" else {"data": payload}
                r = self._session().request(sp.method, sp.action, timeout=25, **kw)
                r.raise_for_status()
                html = r.content.decode(r.encoding or r.apparent_encoding or "utf-8", "replace")
                z = parse_z(html)
                if z is not None:
                    return z
                DEBUG_HTML.write_text(html, encoding="utf-8", errors="replace")
                txt = BeautifulSoup(html, "html.parser").get_text(" ", strip=True)[:400]
                m = re.search(r"(Sorry[^.]*\.)", txt)
                raise ParseFail(f"{m.group(1) if m else 'Website answered but no Z found.'} "
                                f"[sent {sp.method.upper()} {sp.action} {dict(payload)}]")
            except ParseFail:
                raise
            except Exception as e:  # noqa
                if attempt == 2:
                    raise SiteError(str(e))
                time.sleep(0.6 * 2 ** attempt)


class BrowserBackend:
    """Drives the real calculator page in a headless browser (for JavaScript calculators).
    Needs:  pip install playwright   and   playwright install chromium"""
    name = "browser"
    single_thread = True

    def __init__(self, log):
        from playwright.sync_api import sync_playwright
        self.log = log
        self.spec = None
        self.pw = sync_playwright().start()
        self.browser = self.pw.chromium.launch(headless=True)
        self.page = self.browser.new_page()
        diag = []
        for url in SITE_URLS:
            try:
                self.page.goto(url, timeout=30000, wait_until="domcontentloaded")
                self.page.wait_for_timeout(500)
                html = self.page.content()
                DEBUG_HTML.write_text(html, encoding="utf-8", errors="replace")
                spec = discover(html, self.page.url)
                if spec:
                    self.spec = spec
                    log(f"Calculator page opened: {self.page.url}")
                    return
                diag.append(f"{url}: loaded, but no calculator inputs. Elements: {elements_summary(html)}")
            except Exception as e:  # noqa
                diag.append(f"{url}: {str(e)[:100]}")
        self.close()
        raise SiteError("Browser could not find the calculator: " + " | ".join(diag))

    def _read_z(self):
        pg = self.page
        vals = pg.evaluate("""()=>[...document.querySelectorAll('input,textarea')].map(e=>[
            (e.name||'')+' '+(e.id||'')+' '+(e.closest('tr')?e.closest('tr').innerText:''), e.value])""")
        body = pg.inner_text("body").replace("<", " ").replace("&", " ")
        return parse_z("<html><body>" + body + "</body></html>", vals)

    def _set_unit(self, sel, opt):
        pg = self.page
        if sel[3] == "radio":
            pg.check(f"{sel[2]}[value='{opt}']")
        else:
            try:
                pg.select_option(sel[2], value=opt)
            except Exception:
                pg.select_option(sel[2], label=opt)

    def z(self, p, p_unit, t, t_unit):
        sp, pg = self.spec, self.page
        p, popt, t, topt = plan_units(sp, p, p_unit, t, t_unit)
        if popt is not None:
            self._set_unit(sp.p_sel, popt)
        if topt is not None:
            self._set_unit(sp.t_sel, topt)
        pg.fill(sp.t_css, fmtnum(t))
        pg.fill(sp.p_css, fmtnum(p))
        prev = self._read_z()
        btns = pg.query_selector_all("input[type=submit], input[type=button], button")
        pick = next((b for b in btns if re.search(
            r"calc|berech|comput|submit|run|start|ok|go",
            ((b.get_attribute("value") or "") + " " + (b.inner_text() or "")).lower())), btns[0] if btns else None)
        if pick:
            pick.click()
        else:
            pg.press(sp.p_css, "Enter")
        z = None
        for _ in range(14):
            pg.wait_for_timeout(150)
            z = self._read_z()
            if z is not None and z != prev:
                return z
        if z is not None:
            return z
        raise SiteError("Z not found on the page after calculating")

    def close(self):
        try:
            self.browser.close()
            self.pw.stop()
        except Exception:
            pass


def make_backend(log):
    err1 = None
    try:
        be = RequestsBackend(log)
        be.z(10, "bar", 30, "c")   # smoke test: does the plain-HTTP path really return Z?
        return be
    except SiteError as e:
        err1 = str(e)
    log("Plain HTTP mode not possible; switching to a real browser (Playwright)...")
    try:
        be = BrowserBackend(log)
    except ImportError:
        raise SiteError("Needs a browser engine: pip install playwright && playwright install chromium. "
                        "(Plain mode said: " + err1[:600] + ")")
    try:
        be.z(10, "bar", 30, "c")
    except Exception as e:  # noqa
        html = be.page.content()
        be.close()
        raise SiteError(f"Browser opened the page but could not read Z ({e}). Spec: {be.spec.describe()}. "
                        f"Elements: {elements_summary(html)}")
    return be


# --------------------------------------------------------------------------- Excel analysis
P_KEY = re.compile(r"press|psi|bar|kpa|mpa|^pi$|^p$|^p\d*$|pressure|\bp[ _]?(in|i)\b|^\d+(\.\d+)?\s*(psi|bar|kpa|mpa)", re.I)


def analyze_sheet(ws_val, default_unit):
    """Return dict(header_row, p_col, t_col, unit, rows[(row, P)]) or None."""
    max_r, max_c = ws_val.max_row, ws_val.max_column
    header_row = None
    for r in range(1, min(max_r, 15) + 1):
        for c in range(1, max_c + 1):
            v = ws_val.cell(r, c).value
            if isinstance(v, str) and P_KEY.search(v.strip()) and "time" not in v.lower():
                header_row = r
                break
        if header_row:
            break
    if header_row is None:
        header_row = 1
    heads = {c: ws_val.cell(header_row, c).value for c in range(1, max_c + 1)}

    def numeric_count(c):
        n = 0
        for r in range(header_row + 1, max_r + 1):
            v = ws_val.cell(r, c).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                n += 1
        return n

    p_col = None
    for c, h in heads.items():
        if isinstance(h, str) and P_KEY.search(h.strip()) and "time" not in h.lower() and numeric_count(c):
            p_col = c
            break
    if p_col is None:  # first numeric column that is not time
        for c, h in heads.items():
            if not (isinstance(h, str) and "time" in h.lower()) and numeric_count(c):
                p_col = c
                break
    if p_col is None:
        return None
    t_col = next((c for c, h in heads.items() if isinstance(h, str) and "time" in h.lower()), None)
    unit = unit_from_text(heads[p_col]) or unit_from_text(ws_val.title) or default_unit
    rows = []
    for r in range(header_row + 1, max_r + 1):
        v = ws_val.cell(r, p_col).value
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            rows.append((r, float(v)))
    if not rows:
        return None
    return {"header_row": header_row, "p_col": p_col, "t_col": t_col, "unit": unit, "rows": rows, "heads": heads,
            "header": str(heads[p_col]), "last_col": max(c for c, h in heads.items() if h not in (None, ""))}


# --------------------------------------------------------------------------- calculations
TO_PA = {"psi": 6894.757293168, "bar": 1e5, "kpa": 1e3, "mpa": 1e6, "atm": 101325.0, "pa": 1.0}
NOTES_SHEET = "Calculation Notes"
SKIP_SHEETS = {NOTES_SHEET, "Z_settings"}

DEFAULTS = {
    "gas_vol": 0.001,          # m3 (1 L)  - the "0.001" of the original workbook
    "R": 8.314,                # J/(mol K)
    "M": 44.01,                # g/mol CO2 (original workbook used 44)
    "sol_vol_ml": 300.0,       # mL solution
    "conc_factor": 1001.142303,
    "density": 1.0,            # g/mL (1.000 reproduces the original mol x 1000 / 300)
    "sol_mass_g": None,        # optional: overrides volume x density
    "err_mode": "rel",         # "rel" = % of mol/kg, "abs" = fixed mol/kg
    "err_rel": 5.0,            # percent
    "err_abs": 0.001,          # mol/kg
}


def _num(form, key, label, default, positive=True, optional=False):
    raw = str(form.get(key) if form.get(key) is not None else "").strip()
    if raw == "":
        return None if optional else default
    try:
        v = float(raw.replace(",", "."))
    except ValueError:
        raise ValueError(f"{label} must be a number")
    if positive and v <= 0:
        raise ValueError(f"{label} must be greater than 0")
    return v


def parse_cfg(form):
    d = DEFAULTS
    cfg = {
        "temp": _num(form, "temp", "Temperature", 30.0, positive=False),
        "temp_unit": (form.get("temp_unit") or "c").lower(),
        "unit": (form.get("unit") or "psi").lower(),
        "gas_vol": _num(form, "gas_vol", "Gas volume (m3)", d["gas_vol"]),
        "R": _num(form, "R", "Gas constant R", d["R"]),
        "M": _num(form, "M", "CO2 molar mass", d["M"]),
        "sol_vol_ml": _num(form, "sol_vol_ml", "Solution volume (mL)", d["sol_vol_ml"]),
        "conc_factor": _num(form, "conc_factor", "Concentration factor", d["conc_factor"]),
        "density": _num(form, "density", "Solution density", d["density"]),
        "sol_mass_g": _num(form, "sol_mass_g", "Solution mass (g)", None, optional=True),
        "err_mode": (form.get("err_mode") or "rel").lower(),
        "err_rel": _num(form, "err_rel", "Error percent", d["err_rel"], positive=False),
        "err_abs": _num(form, "err_abs", "Fixed error", d["err_abs"], positive=False),
    }
    if cfg["unit"] not in TO_PA:
        raise ValueError("Unknown pressure unit")
    if cfg["temp_unit"] not in ("c", "k"):
        raise ValueError("Temperature unit must be C or K")
    if t_kelvin(cfg) <= 0:
        raise ValueError("Temperature must be above absolute zero")
    if cfg["err_mode"] not in ("rel", "abs"):
        raise ValueError("Error mode must be 'rel' or 'abs'")
    return cfg


def t_kelvin(cfg):
    """T(K) = T(C) + 273.15"""
    return cfg["temp"] + 273.15 if cfg["temp_unit"] == "c" else cfg["temp"]


def solution_mass_kg(cfg):
    if cfg["sol_mass_g"]:
        return cfg["sol_mass_g"] / 1000.0
    return cfg["sol_vol_ml"] * cfg["density"] / 1000.0      # mL x g/mL = g ; /1000 = kg


def calc_row(cfg, unit, p1, p2, z1, z2):
    """All derived quantities for one row. Returns None if Z1 or Z2 is missing."""
    if z1 is None or z2 is None:
        return None
    pi = p1 * TO_PA[unit]                                     # Pa
    pf = p2 * TO_PA[unit]                                     # Pa
    mol = (pi / z1 - pf / z2) * cfg["gas_vol"] / (cfg["R"] * t_kelvin(cfg))
    g = mol * cfg["M"]
    gl = g / (cfg["sol_vol_ml"] / 1000.0)
    conc = gl * cfg["conc_factor"]
    molkg = mol / solution_mass_kg(cfg)
    err = abs(molkg) * cfg["err_rel"] / 100.0 if cfg["err_mode"] == "rel" else cfg["err_abs"]
    return {"pi": pi, "pf": pf, "mol": mol, "g": g, "gL": gl, "conc": conc, "molkg": molkg, "err": err}


TIME_TOK = re.compile(r"(\d+(?:[.,]\d+)?)\s*(hours?|hrs?|h|minutes?|mins?|m|seconds?|secs?|s)\b", re.I)


def parse_seconds(v):
    """'0 sec' -> 0, '3 min' -> 180, '27 min 30 sec' -> 1650. Plain numbers are taken as seconds."""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip().lower()
    toks = TIME_TOK.findall(s)
    if toks:
        rest = TIME_TOK.sub("", s)
        if re.sub(r"[\s,]|and", "", rest) == "":
            tot = 0.0
            for num, u in toks:
                n = float(num.replace(",", "."))
                tot += n * (3600 if u[0] == "h" else 60 if u[0] == "m" else 1)
            return tot
        return None
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return None


# --------------------------------------------------------------------------- jobs
JOBS = {}


class Job:
    def __init__(self, data, filename, cfg, workers):
        self.id = uuid.uuid4().hex[:10]
        self.data, self.filename, self.cfg, self.workers = data, filename, cfg, workers
        self.events = []
        self.lock = threading.Lock()
        self.cancel = False
        self.done = False
        self.out_path = None
        self.plans = {}
        self.results = {}      # sheet -> {"z1","z1_err","z":{p:Z},"err":{p:msg}}  (kept across retries)
        self.site = None
        self.total = self.finished = self.failed = 0
        self.t0 = time.time()

    def emit(self, **ev):
        with self.lock:
            self.events.append(ev)

    def log(self, msg):
        self.emit(t="log", msg=msg)


def analyse(job):
    wb = openpyxl.load_workbook(io.BytesIO(job.data), data_only=True)
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        info = analyze_sheet(ws, job.cfg["unit"])
        if info is None:
            job.log(f"Sheet '{ws.title}': no pressure column found - skipped")
            continue
        heads = info["heads"]
        sec_col = next((c for c, h in heads.items() if norm(str(h)) == "seconds"), None)
        src = sec_col or info["t_col"]
        info["secs"] = {r: parse_seconds(ws.cell(r, src).value) if src else None for r, _ in info["rows"]}
        info["time_txt"] = {r: (ws.cell(r, info["t_col"]).value if info["t_col"] else None) for r, _ in info["rows"]}
        info["rowp"] = {r: p for r, p in info["rows"]}
        info["has_seconds_col"] = sec_col is not None
        job.plans[ws.title] = info
        job.results[ws.title] = {"z1": None, "z1_err": None, "z": {}, "err": {}}
        job.log(f"Sheet '{ws.title}': pressure column '{info['header']}' unit {UNIT_LABEL[info['unit']]}, "
                f"{len(info['rows'])} rows")
    if not job.plans:
        raise SiteError("No sheet with a numeric pressure column was found.")
    meta = []
    for name, info in job.plans.items():
        meta.append({"sheet": name, "unit": UNIT_LABEL[info["unit"]], "p1": info["rows"][0][1],
                     "rows": [{"r": r, "t": info["time_txt"][r], "p": p,
                               "min": (info["secs"][r] / 60 if info["secs"][r] is not None else None)}
                              for r, p in info["rows"]]})
    job.emit(t="init", sheets=meta, temp=job.cfg["temp"], temp_unit=job.cfg["temp_unit"].upper())


def row_state(job, name, r):
    info, res = job.plans[name], job.results[name]
    p = info["rowp"][r]
    z1, z2, err = res["z1"], res["z"].get(p), res["err"].get(p)
    calc = calc_row(job.cfg, info["unit"], info["rows"][0][1], p, z1, z2)
    if err is None and z2 is None:
        err = "not calculated"
    if err is None and z2 is not None and z1 is None:
        err = "Z1 missing: " + str(res["z1_err"] or "not calculated")
    return {"sheet": name, "row": r, "z2": z2, "err": err, "calc": calc}


def failed_rows(job):
    out = []
    for name, info in job.plans.items():
        for r, p in info["rows"]:
            st = row_state(job, name, r)
            if st["calc"] is None:
                out.append((name, r, p, st["err"] or "not calculated"))
    return out


def run_pass(job):
    """One pass: look up every Z that is still missing (everything on the first pass, only failures on a retry)."""
    cfg = job.cfg
    job.total = job.finished = job.failed = 0
    for name, info in job.plans.items():
        res = job.results[name]
        job.total += (res["z1"] is None) + sum(1 for r, p in info["rows"] if p not in res["z"])
    if job.total == 0:
        return
    job.log("Connecting to the Peace Software website...")
    backend = make_backend(job.log)
    try:
        job.site = getattr(backend.spec, "action", None) or getattr(backend.spec, "page_url", None)
        job.log(f"Connected ({backend.name} mode). Calculating...")
        workers = 1 if getattr(backend, "single_thread", False) else job.workers
        cache, clock = {}, threading.Lock()

        def get_z(p, unit):
            key = (unit, round(p, 10))
            with clock:
                if key in cache:
                    return cache[key]
            last = None
            for attempt in range(3 if backend.name == "browser" else 1):   # requests backend retries itself
                try:
                    z = backend.z(p, unit, cfg["temp"], cfg["temp_unit"])
                    with clock:
                        cache[key] = z
                    return z
                except Exception as e:  # noqa
                    last = e
                    time.sleep(0.6 * 2 ** attempt)
            raise last

        for name, info in job.plans.items():
            if job.cancel:
                break
            res, unit, rows = job.results[name], info["unit"], info["rows"]
            if res["z1"] is None:                                    # Z1: once per sheet
                try:
                    res["z1"], res["z1_err"] = get_z(rows[0][1], unit), None
                except Exception as e:  # noqa
                    res["z1_err"] = str(e)
                    job.log(f"Z1 failed for '{name}': {str(e)[:200]}")
                job.finished += 1
                job.emit(t="z1", sheet=name, z=res["z1"], err=res["z1_err"])
            rows_by_p = {}
            for r, p in rows:
                rows_by_p.setdefault(p, []).append(r)
            missing = [p for p in rows_by_p if p not in res["z"]]

            def work(p):
                if job.cancel:
                    return p, None, "stopped by user"
                try:
                    return p, get_z(p, unit), None
                except Exception as e:  # noqa
                    return p, None, str(e)

            def record(p, z, err):
                if z is not None:
                    res["z"][p] = z
                    res["err"].pop(p, None)
                else:
                    res["err"][p] = err
                for r in rows_by_p[p]:
                    job.finished += 1
                    if z is None:
                        job.failed += 1
                    job.emit(t="row", **row_state(job, name, r))

            if getattr(backend, "single_thread", False):
                for p in missing:
                    record(*work(p))
            else:
                with ThreadPoolExecutor(max_workers=workers) as ex:
                    for res_ in ex.map(work, missing):
                        record(*res_)
    finally:
        if hasattr(backend, "close"):
            backend.close()


def pass_and_write(job):
    try:
        run_pass(job)
    except Exception as e:  # noqa
        job.emit(t="error", msg=str(e))
        traceback.print_exc()
    finally:
        try:
            have = any(r["z1"] is not None or r["z"] for r in job.results.values())
            bad = failed_rows(job)
            job.failed = len(bad)
            sync = {n: {"z1": job.results[n]["z1"], "z1err": job.results[n]["z1_err"],
                        "rows": [row_state(job, n, r) for r, _ in info["rows"]]} for n, info in job.plans.items()}
            job.emit(t="sync", sheets=sync)
            complete = (not bad) and not job.cancel
            if have:
                old = job.out_path
                job.out_path = write_xlsx(job, complete)
                if old and old != job.out_path and Path(old).exists():
                    Path(old).unlink()
            job.emit(t="done", file=(job.out_path.name if job.out_path else None), failed=len(bad),
                     complete=complete, stopped=job.cancel)
        except Exception as e:  # noqa
            job.emit(t="error", msg=f"Could not write the Excel file: {e}")
            traceback.print_exc()
        finally:
            job.done = True


def run_job(job):
    try:
        analyse(job)
    except Exception as e:  # noqa
        job.emit(t="error", msg=str(e))
        job.done = True
        return
    pass_and_write(job)


# --------------------------------------------------------------------------- Excel output
def out_columns(unit_label):
    """(key, header, aliases, number_format)"""
    N = norm
    return [
        ("seconds", "Seconds", {N("Seconds")}, "0.###"),
        ("minutes", "Minutes", {N("Minutes")}, "0.0000"),
        ("p1", f"P1 ({unit_label})", {N("P1"), N(f"P1 ({unit_label})")}, "0.0##"),
        ("p2", f"P2 ({unit_label})", {N("P2"), N(f"P2 ({unit_label})")}, "0.0##"),
        ("z1", "Z1", {N("Z1")}, "0.00000"),
        ("z2", "Z2", {N("Z2")}, "0.00000"),
        ("pi", "Pi (Pa)", {N("Pi (Pa)")}, "#,##0.0"),
        ("pf", "Pf (Pa)", {N("Pf (Pa)")}, "#,##0.0"),
        ("mol", "Absorbed mol", {N("Absorbed mol"), N("Absorbed moles")}, "0.000000"),
        ("g", "Absorbed CO2 (gm)", {N("Absorbed CO2 (gm)"), N("Absorbed CO2 (g)")}, "0.0000"),
        ("gL", "Quantity (gm/L)", {N("Quantity (gm/L)"), N("Quantity (g/L)")}, "0.0000"),
        ("conc", "Concentration", {N("Concentration")}, "0.00"),
        ("molkg", "mol/kg", {N("mol/kg")}, "0.000000"),
        ("err", "Error", {N("Error")}, "0.000000"),
    ]


def write_xlsx(job, complete):
    from copy import copy
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    cfg = job.cfg
    wb = openpyxl.load_workbook(io.BytesIO(job.data))        # fresh copy: the uploaded bytes are never modified
    for junk in SKIP_SHEETS:
        if junk in wb.sheetnames:
            del wb[junk]
    for name, info in job.plans.items():
        ws, res = wb[name], job.results[name]
        hr, heads = info["header_row"], info["heads"]
        reserved = {c for c in (info["p_col"], info["t_col"]) if c}
        existing = {}
        for c, h in heads.items():
            if c not in reserved and h not in (None, ""):
                existing.setdefault(norm(str(h)), c)
        nxt, used, colmap, created = info["last_col"] + 1, set(), {}, set()
        src_hdr = ws.cell(hr, info["p_col"])
        p1 = info["rows"][0][1]
        for key, header, aliases, fmt in out_columns(UNIT_LABEL[info["unit"]]):
            col = next((existing[a] for a in aliases if a in existing and existing[a] not in used), None)
            if col is None:
                col = nxt
                nxt += 1
                created.add(key)
                cell = ws.cell(hr, col, header)
                cell._style = copy(src_hdr._style)
                cell.font = Font(name=src_hdr.font.name, size=src_hdr.font.size, bold=True)
                ws.column_dimensions[get_column_letter(col)].width = max(12, len(header) + 3)
            used.add(col)
            colmap[key] = (col, fmt)
        for r, p in info["rows"]:
            st = row_state(job, name, r)
            c = st["calc"] or {}
            sec = info["secs"][r]
            vals = {"seconds": sec, "minutes": (sec / 60.0 if sec is not None else None), "p1": p1, "p2": p,
                    "z1": res["z1"], "z2": st["z2"], "pi": p1 * TO_PA[info["unit"]], "pf": p * TO_PA[info["unit"]],
                    "mol": c.get("mol"), "g": c.get("g"), "gL": c.get("gL"), "conc": c.get("conc"),
                    "molkg": c.get("molkg"), "err": c.get("err")}
            for key, (col, fmt) in colmap.items():
                if key == "seconds" and info["has_seconds_col"]:
                    continue                                   # keep the original Seconds column untouched
                cell = ws.cell(r, col)
                cell.value = vals[key]                         # numbers, or empty if not calculated
                cell.number_format = fmt
    build_notes(job, wb, complete)
    out = OUT_DIR / (Path(job.filename).stem + ("_Z.xlsx" if complete else "_Z_INCOMPLETE.xlsx"))
    wb.save(out)
    return out


def build_notes(job, wb, complete):
    from openpyxl.styles import Alignment, Font, PatternFill
    cfg = job.cfg
    ws = wb.create_sheet(NOTES_SHEET)
    bad = failed_rows(job)
    rows = []

    def add(*cells, bold=False, fill=None):
        rows.append((cells, bold, fill))

    nrows = sum(len(i["rows"]) for i in job.plans.values())
    if complete:
        add("STATUS", f"COMPLETE - all {nrows} rows calculated", bold=True, fill="C6EFCE")
    else:
        why = "stopped by user; " if job.cancel else ""
        add("STATUS", f"INCOMPLETE - {why}{len(bad)} of {nrows} rows have no result (listed at the bottom). "
                      "Do not use the blank rows as final data.", bold=True, fill="FFC7CE")
    add("Generated", time.strftime("%Y-%m-%d %H:%M:%S"))
    add("Source file", job.filename, None, "The uploaded workbook was not modified; all original sheets and cells are kept.")
    add("Z source", "Peace Software online CO2 calculator", None,
        f"{job.site or ''} - NOT CoolProp, no estimates. A failed website request leaves the cell blank.")
    add()
    add("SETTINGS USED", "Value", "Unit", "Note", bold=True, fill="D9E1F2")
    add("Temperature", cfg["temp"], "deg C" if cfg["temp_unit"] == "c" else "K",
        "Used for the website lookups AND for T(K) in the moles equation.")
    add("T (K)", t_kelvin(cfg), "K", "T(K) = T(C) + 273.15")
    add("Gas volume V", cfg["gas_vol"], "m3", "The original '0.001'. Only dimensionally correct if it is a volume in m3 "
        "(0.001 m3 = 1 L). CONFIRM this is the gas volume of your cell.")
    add("Gas constant R", cfg["R"], "J/(mol K)", "Same as the original workbook.")
    add("CO2 molar mass M", cfg["M"], "g/mol", "Original workbook used 44. 44.01 is used here (difference +0.023 %).")
    add("Solution volume", cfg["sol_vol_ml"], "mL", "Original workbook: 300.")
    add("Concentration factor", cfg["conc_factor"], "(not documented)",
        "Taken unchanged from the original workbook; its units and source are NOT confirmed. Observation only: "
        "1000/1001.142303 = 0.998859, i.e. it looks like a g/L -> mg/kg conversion with a density of about "
        "0.99886 g/mL (water near 16-17 deg C). Not verified; a CaCl2 solution is denser than water.")
    add("Solution density", cfg["density"], "g/mL", "Used for mol/kg. 1.000 reproduces the original 'mol x 1000 / 300'. "
        "CaCl2 / NaCl / MgCl2 brines are denser than 1.000 - set the true density for a correct mol/kg.")
    add("Solution mass (override)", cfg["sol_mass_g"] if cfg["sol_mass_g"] else "not set", "g",
        "If set, this mass is used for mol/kg instead of volume x density.")
    add("Solution mass used", solution_mass_kg(cfg), "kg", "= solution volume (mL) x density (g/mL) / 1000, or the override")
    if cfg["err_mode"] == "rel":
        add("Error definition", f"{cfg['err_rel']} % of |mol/kg|", "mol/kg", "As requested: 5 % error of the mol/kg value.")
    else:
        add("Error definition", f"fixed {cfg['err_abs']}", "mol/kg", "As requested: fixed error of 0.001 mol/kg.")
    add()
    add("FORMULAS", "Formula", "Unit", "Note", bold=True, fill="D9E1F2")
    add("Seconds", "parsed from the Time text ('3 min' = 180, '27 min 30 sec' = 1650)", "s",
        "Time text is never changed. Plain numbers are read as seconds.")
    add("Minutes", "Seconds / 60", "min")
    add("P1, Z1", "first pressure of the sheet, repeated; Z1 looked up once per sheet", "-")
    add("Pi (Pa)", "P1 x 6894.757293168 (psi -> Pa)", "Pa", "other units: bar x 1e5, kPa x 1e3, MPa x 1e6")
    add("Pf (Pa)", "P2 x 6894.757293168", "Pa")
    add("Absorbed mol", "((Pi / Z1) - (Pf / Z2)) x V / (R x T)", "mol",
        "Check: Pa x m3 = J; J / (J mol-1 K-1 x K) = mol. Same equation as in your specification; the 0.001 is V in m3.")
    add("Absorbed CO2 (gm)", "Absorbed mol x M", "g")
    add("Quantity (gm/L)", "Absorbed CO2 (g) / (solution volume mL / 1000)", "g/L", "Original: x 1000 / 300 (same at 300 mL)")
    add("Concentration", "Quantity (gm/L) x concentration factor", "not confirmed", "See the factor above.")
    add("mol/kg", "Absorbed mol / solution mass (kg)", "mol/kg",
        "The original 'mol x 1000 / 300' is mol per litre of solution; it equals mol/kg only if the solution is 1.000 g/mL.")
    add("Error", "5 % of |mol/kg|  (or fixed value, see settings)", "mol/kg")
    add()
    add("CHECKS", bold=True, fill="D9E1F2")
    add("Template formulas", "The uploaded CaCl2 template had no formulas, only Time and pressure; formulas follow your written specification.")
    add("Blank cells", "Z1/Z2 and everything that depends on them stay empty if the website lookup failed. Pi/Pf/Seconds/Minutes are always filled.")
    add("Pressure limit", "The website only calculates pressures >= 1 bar (14.504 psi); lower values are reported as failed rows.")
    add()
    add("PER SHEET", "Z1", "P1 / unit", "Rows", bold=True, fill="D9E1F2")
    for name, info in job.plans.items():
        res = job.results[name]
        unparsed = [r for r, _ in info["rows"] if info["secs"][r] is None]
        note = f"{len(info['rows'])} rows" + (f"; Time not readable in rows {unparsed}" if unparsed else "")
        add(name, res["z1"] if res["z1"] is not None else f"FAILED: {res['z1_err']}",
            f"{info['rows'][0][1]} {UNIT_LABEL[info['unit']]}", note)
    if bad:
        add()
        add("FAILED / MISSING ROWS", "Sheet row", "Pressure", "Error message", bold=True, fill="FFC7CE")
        for name, r, p, err in bad:
            add(name, r, p, str(err)[:500])
    for cells, bold, fill in rows:
        ws.append(list(cells))
        rr = ws.max_row
        for c in ws[rr]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if bold:
                c.font = Font(bold=True)
            if fill and c.column <= 4:
                c.fill = PatternFill("solid", fgColor=fill)
    for col, w in zip("ABCD", (28, 44, 22, 90)):
        ws.column_dimensions[col].width = w


# --------------------------------------------------------------------------- web app
app = Flask(__name__)


@app.get("/")
def index():
    return Response(PAGE, mimetype="text/html")


@app.post("/start")
def start():
    f = request.files.get("file")
    if not f:
        return jsonify(error="No file uploaded"), 400
    try:
        cfg = parse_cfg(request.form)
        workers = max(1, min(8, int(request.form.get("workers", "4"))))
    except ValueError as e:
        return jsonify(error=str(e)), 400
    job = Job(f.read(), f.filename or "workbook.xlsx", cfg, workers)
    JOBS[job.id] = job
    threading.Thread(target=run_job, args=(job,), daemon=True).start()
    return jsonify(id=job.id)


@app.post("/retry/<jid>")
def retry(jid):
    job = JOBS.get(jid)
    if not job or not job.done:
        return jsonify(error="nothing to retry"), 400
    job.done, job.cancel, job.t0 = False, False, time.time()
    threading.Thread(target=pass_and_write, args=(job,), daemon=True).start()
    return jsonify(ok=True)


@app.get("/events/<jid>")
def events(jid):
    job = JOBS.get(jid)
    if not job:
        return jsonify(error="unknown job"), 404
    after = int(request.args.get("after", 0))
    with job.lock:
        ev = job.events[after:after + 3000]
    el = time.time() - job.t0
    rate = job.finished / el if el > 0 else 0
    eta = (job.total - job.finished) / rate if rate > 0 else None
    return jsonify(events=ev, next=after + len(ev), done=job.done, total=job.total,
                   finished=job.finished, failed=job.failed, elapsed=el, eta=eta)


@app.post("/cancel/<jid>")
def cancel(jid):
    if jid in JOBS:
        JOBS[jid].cancel = True
    return jsonify(ok=True)


@app.get("/download/<jid>")
def download(jid):
    job = JOBS.get(jid)
    if not job or not job.out_path or not Path(job.out_path).exists():
        return "Not ready", 404
    return send_file(job.out_path, as_attachment=True, download_name=Path(job.out_path).name)


@app.get("/probe")
def probe():
    """Connection test: shows what the app found on the website and one Z value."""
    msgs = []
    try:
        be = make_backend(msgs.append)
        unit = request.args.get("unit", "bar")
        val = {"psi": 100, "bar": 10, "kpa": 1000, "mpa": 1}.get(unit, 10)     # the site needs >= 1 bar
        z = be.z(val, unit, 30, "c")
        out = {"ok": True, "mode": be.name, "form": be.spec.describe(), "test": f"Z at {val} {unit}, 30 C = {z}"}
        if hasattr(be, "close"):
            be.close()
        return jsonify(out)
    except Exception as e:  # noqa
        return jsonify(ok=False, error=str(e), log=msgs, hint=f"Page saved to {DEBUG_HTML.name} for inspection")


PAGE = r"""<!doctype html><html><head><meta charset="utf-8"><title>CO2 absorption dashboard</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
<style>
:root{--bg:#0f1419;--card:#18202a;--line:#2a3441;--tx:#e6edf3;--mut:#8b98a5;--ac:#4cc2ff;--ok:#3fb950;--bad:#f85149}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--tx);font:14px system-ui,Segoe UI,sans-serif}
.wrap{max-width:1250px;margin:0 auto;padding:20px}h1{font-size:20px;margin:0 0 4px}.sub{color:var(--mut);margin-bottom:16px}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px;margin-bottom:14px}
.row{display:flex;flex-wrap:wrap;gap:12px;align-items:end}label{display:flex;flex-direction:column;gap:4px;color:var(--mut);font-size:12px}
input,select{background:#0f1419;color:var(--tx);border:1px solid var(--line);border-radius:6px;padding:7px 9px;font-size:14px}
button{background:var(--ac);color:#04121c;border:0;border-radius:6px;padding:9px 14px;font-weight:600;cursor:pointer}
button.sec{background:transparent;color:var(--tx);border:1px solid var(--line)}button:disabled{opacity:.4;cursor:not-allowed}
summary{cursor:pointer;color:var(--ac);font-weight:600}details .row{margin-top:12px}.hint{color:var(--mut);font-size:12px;margin-top:8px}
.bar{height:10px;background:#0f1419;border-radius:6px;overflow:hidden;border:1px solid var(--line)}.bar>div{height:100%;width:0;background:var(--ac);transition:width .2s}
.stats{display:flex;gap:18px;color:var(--mut);margin-top:8px;flex-wrap:wrap}.stats b{color:var(--tx)}
.tabs{display:flex;gap:6px;margin-bottom:8px;flex-wrap:wrap}.tab{padding:5px 11px;border:1px solid var(--line);border-radius:6px;cursor:pointer;color:var(--mut)}.tab.on{color:var(--tx);border-color:var(--ac)}
.tw{max-height:420px;overflow:auto;border:1px solid var(--line);border-radius:6px}table{border-collapse:collapse;width:100%;font-variant-numeric:tabular-nums;white-space:nowrap}
th,td{padding:5px 10px;text-align:right;border-bottom:1px solid var(--line)}th{position:sticky;top:0;background:var(--card);color:var(--mut)}
td.z{color:var(--ok)}td.e{color:var(--bad)}td.w{color:var(--mut)}
#msg.bad{color:var(--bad);font-weight:600}#msg.ok{color:var(--ok);font-weight:600}
#fail td{text-align:left;white-space:normal}
#log{height:120px;overflow:auto;font:12px ui-monospace,monospace;color:var(--mut);white-space:pre-wrap}
#probe{white-space:pre-wrap;font:12px ui-monospace,monospace;color:var(--mut);margin-top:8px}
</style></head><body><div class="wrap">
<h1>CO&#8322; absorption dashboard</h1><div class="sub">Z values come from the Peace Software online CO&#8322; calculator &mdash; each pressure is entered into the website automatically, then the absorption results are calculated and written to Excel.</div>
<div class="card"><div class="row">
<label>Excel workbook<input type="file" id="file" accept=".xlsx,.xlsm"></label>
<label>Temperature<input type="number" id="temp" value="30" step="any" style="width:90px"></label>
<label>Unit<select id="tu"><option value="c">&deg;C</option><option value="k">K</option></select></label>
<label>Pressure unit if not in header<select id="unit"><option>psi</option><option>bar</option><option>kpa</option><option>mpa</option></select></label>
<label>Parallel requests<input type="number" id="workers" value="4" min="1" max="8" style="width:70px"></label>
<button id="go">Start</button><button class="sec" id="stop" disabled>Stop</button><button class="sec" id="pr">Test website connection</button>
</div>
<details style="margin-top:12px"><summary>Calculation settings (defaults follow your workbook)</summary><div class="row">
<label>Gas volume V (m&sup3;)<input id="gas_vol" value="0.001" style="width:100px"></label>
<label>R (J/mol/K)<input id="R" value="8.314" style="width:80px"></label>
<label>CO&#8322; molar mass (g/mol)<input id="M" value="44.01" style="width:90px"></label>
<label>Solution volume (mL)<input id="sol_vol_ml" value="300" style="width:90px"></label>
<label>Concentration factor<input id="conc_factor" value="1001.142303" style="width:120px"></label>
<label>Solution density (g/mL)<input id="density" value="1.000" style="width:90px"></label>
<label>Solution mass (g) - optional<input id="sol_mass_g" value="" placeholder="auto" style="width:100px"></label>
<label>Error mode<select id="err_mode"><option value="rel">% of mol/kg</option><option value="abs">fixed mol/kg</option></select></label>
<label>Error %<input id="err_rel" value="5" style="width:60px"></label>
<label>Fixed error (mol/kg)<input id="err_abs" value="0.001" style="width:80px"></label>
</div><div class="hint">V = 0.001 only works as 0.001 m&sup3; (1 L). Density 1.000 reproduces your original mol/kg; set the real brine density for a true mol/kg.</div></details>
<div id="probe"></div></div>
<div class="card"><div class="bar"><div id="fill"></div></div>
<div class="stats"><span>Done <b id="sd">0</b>/<b id="st">0</b></span><span>Failed <b id="sf">0</b></span><span>Elapsed <b id="se">0s</b></span><span>ETA <b id="sa">-</b></span><span id="msg"></span></div>
<div style="margin-top:10px;display:flex;gap:10px"><a id="dl" style="display:none"><button>Download Excel</button></a><button class="sec" id="retry" style="display:none">Retry failed rows only</button></div></div>
<div class="card"><div class="tabs" id="tabs"></div><div class="tw"><table><thead><tr id="thr"></tr></thead><tbody id="tb"></tbody></table></div></div>
<div class="card" id="failcard" style="display:none"><b style="color:var(--bad)">Rows without a result</b><div class="tw" style="margin-top:8px;max-height:220px"><table id="fail"></table></div></div>
<div class="card"><div id="log"></div></div></div>
<script>
const $=id=>document.getElementById(id);let jid=null,after=0,timer=null,S={},names=[],cur=0;
const f=(n,d)=>n==null?'':Number(n).toFixed(d);
const esc=s=>String(s==null?'':s).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/"/g,'&quot;');
const td=(v,d)=>'<td class="'+(v==null?'w':'z')+'">'+(v==null?'':f(v,d))+'</td>';
function rowHtml(s,r){const c=r.calc||{};
 const z1=s.z1==null?'<td class="'+(s.z1err?'e':'w')+'" title="'+esc(s.z1err)+'">'+(s.z1err?'ERR':'...')+'</td>':'<td class="z">'+f(s.z1,5)+'</td>';
 const z2='<td class="'+(r.err?'e':r.z2==null?'w':'z')+'" title="'+esc(r.err)+'">'+(r.err?'ERR':r.z2==null?'...':f(r.z2,5))+'</td>';
 return '<td>'+r.r+'</td><td>'+esc(r.t)+'</td><td>'+f(r.min,3)+'</td><td>'+s.p1+'</td><td>'+r.p+'</td>'+z1+z2+td(c.pi,1)+td(c.pf,1)+td(c.mol,6)+td(c.g,4)+td(c.gL,4)+td(c.conc,2)+td(c.molkg,6)+td(c.err,6)}
function render(){const s=S[names[cur]];if(!s)return;
 $('thr').innerHTML=['Row','Time','Min','P1 ('+s.unit+')','P2 ('+s.unit+')','Z1','Z2','Pi (Pa)','Pf (Pa)','Abs. mol','CO2 (g)','g/L','Conc.','mol/kg','Error'].map(h=>'<th>'+h+'</th>').join('');
 $('tb').innerHTML=s.rows.map((r,i)=>'<tr id="r'+i+'">'+rowHtml(s,r)+'</tr>').join('')}
function tabs(){$('tabs').innerHTML=names.map((k,i)=>'<div class="tab '+(i==cur?'on':'')+'" onclick="cur='+i+';tabs();render()">'+esc(k)+'</div>').join('')}
function fails(){let h='<tr><th>Sheet</th><th>Row</th><th>Pressure</th><th>Error message</th></tr>',n=0;
 for(const k of names){const s=S[k];for(const r of s.rows){if(!r.calc){n++;h+='<tr><td>'+esc(k)+'</td><td>'+r.r+'</td><td>'+r.p+' '+s.unit+'</td><td>'+esc(r.err||s.z1err||'not calculated')+'</td></tr>'}}}
 $('fail').innerHTML=h;$('failcard').style.display=n?'block':'none'}
function log(m){const l=$('log');l.textContent+=m+'\n';l.scrollTop=l.scrollHeight}
function dur(s){s=Math.round(s);return s>=60?Math.floor(s/60)+'m '+(s%60)+'s':s+'s'}
function setMsg(t,c){$('msg').textContent=t;$('msg').className=c||''}
function apply(ev){for(const e of ev){
 if(e.t=='init'){S={};names=[];for(const m of e.sheets){const idx={};m.rows.forEach((r,i)=>idx[r.r]=i);names.push(m.sheet);S[m.sheet]={unit:m.unit,p1:m.p1,z1:null,z1err:null,idx,rows:m.rows.map(r=>({r:r.r,t:r.t,p:r.p,min:r.min,z2:null,err:null,calc:null}))}}cur=0;tabs();render()}
 else if(e.t=='z1'){const s=S[e.sheet];s.z1=e.z;s.z1err=e.err;if(names[cur]==e.sheet)render()}
 else if(e.t=='row'){const s=S[e.sheet],i=s.idx[e.row];Object.assign(s.rows[i],{z2:e.z2,err:e.err,calc:e.calc});
   if(names[cur]==e.sheet){const tr=$('r'+i);if(tr){tr.innerHTML=rowHtml(s,s.rows[i]);tr.scrollIntoView({block:'nearest'})}}}
 else if(e.t=='sync'){for(const k in e.sheets){const s=S[k],d=e.sheets[k];s.z1=d.z1;s.z1err=d.z1err;for(const r of d.rows){const i=s.idx[r.row];Object.assign(s.rows[i],{z2:r.z2,err:r.err,calc:r.calc})}}render();fails()}
 else if(e.t=='log')log(e.msg);
 else if(e.t=='error'){log('ERROR: '+e.msg);setMsg('Error: '+e.msg,'bad')}
 else if(e.t=='done'){if(e.file){$('dl').href='/download/'+jid;$('dl').style.display='inline-block'}
   $('retry').style.display=e.complete?'none':'inline-block';
   if(e.complete)setMsg('Finished - all rows calculated','ok');
   else setMsg((e.stopped?'Stopped. ':'')+'INCOMPLETE: '+e.failed+' rows have no result. The workbook is marked INCOMPLETE.','bad')}}}
async function poll(){const r=await (await fetch('/events/'+jid+'?after='+after)).json();after=r.next;apply(r.events);
 $('sd').textContent=r.finished;$('st').textContent=r.total;$('sf').textContent=r.failed;$('se').textContent=dur(r.elapsed);$('sa').textContent=r.eta==null||r.done?'-':dur(r.eta);
 $('fill').style.width=(r.total?100*r.finished/r.total:0)+'%';
 if(r.done&&r.events.length==0){clearInterval(timer);$('go').disabled=false;$('stop').disabled=true}}
const IDS=['gas_vol','R','M','sol_vol_ml','conc_factor','density','sol_mass_g','err_mode','err_rel','err_abs'];
$('go').onclick=async()=>{const file=$('file').files[0];if(!file){alert('Choose an Excel file first');return}
 const fd=new FormData();fd.append('file',file);fd.append('temp',$('temp').value);fd.append('temp_unit',$('tu').value);fd.append('unit',$('unit').value);fd.append('workers',$('workers').value);
 for(const k of IDS)fd.append(k,$(k).value);
 $('go').disabled=true;$('stop').disabled=false;$('dl').style.display='none';$('retry').style.display='none';$('failcard').style.display='none';$('log').textContent='';setMsg('');after=0;
 const r=await (await fetch('/start',{method:'POST',body:fd})).json();if(r.error){alert(r.error);$('go').disabled=false;$('stop').disabled=true;return}
 jid=r.id;timer=setInterval(poll,400)};
$('stop').onclick=()=>fetch('/cancel/'+jid,{method:'POST'});
$('retry').onclick=async()=>{const r=await (await fetch('/retry/'+jid,{method:'POST'})).json();if(r.error){alert(r.error);return}
 $('retry').style.display='none';$('dl').style.display='none';$('go').disabled=true;$('stop').disabled=false;setMsg('Retrying failed rows...');timer=setInterval(poll,400)};
$('pr').onclick=async()=>{$('probe').textContent='Testing...';const r=await (await fetch('/probe?unit='+$('unit').value)).json();$('probe').textContent=JSON.stringify(r,null,2)};
</script></body></html>"""


if __name__ == "__main__":
    threading.Timer(1.0, lambda: webbrowser.open(f"http://127.0.0.1:{PORT}")).start()
    print(f"Dashboard running at http://127.0.0.1:{PORT}  (Ctrl+C to stop)")
    app.run(host="127.0.0.1", port=PORT, threaded=True)
